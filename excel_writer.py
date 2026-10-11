"""Write extracted invoice rows to a formatted Excel workbook."""

from __future__ import annotations

from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from extractor import EXTRACTION_METHOD_COLUMN, FIELD_NAMES

SOURCE_COLUMN = "Source File"
EXCEL_COLUMNS = [*FIELD_NAMES, SOURCE_COLUMN, EXTRACTION_METHOD_COLUMN]


def write_excel(rows: list[dict[str, Any]], output_path: str | Path) -> Path:
    """Create a clean, consistently ordered Invoices worksheet."""
    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    dataframe = pd.DataFrame(rows, columns=EXCEL_COLUMNS)

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        dataframe.to_excel(writer, index=False, sheet_name="Invoices")
        worksheet = writer.sheets["Invoices"]
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions

        for cell in worksheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(fill_type="solid", fgColor="1F4E78")

        for column_index, column_cells in enumerate(worksheet.columns, start=1):
            max_length = max(
                (len(str(cell.value)) if cell.value is not None else 0)
                for cell in column_cells
            )
            worksheet.column_dimensions[get_column_letter(column_index)].width = min(
                max(max_length + 2, 14), 40
            )

    return path
